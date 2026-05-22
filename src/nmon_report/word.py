from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches

from .analyzer import build_server_overview
from .models import ServerReport

THEME_TITLES = {
    "tps": "TPS使用情况",
    "cpu": "CPU使用情况",
    "memory": "内存使用情况",
    "disk": "磁盘使用情况",
    "network": "网络使用情况",
    "filesystem": "文件系统使用情况",
    "system": "系统使用情况",
    "process": "进程使用情况",
    "adapter": "适配器使用情况",
    "jfs": "JFS使用情况",
    "lpar": "LPAR使用情况",
    "unclassified": "其他指标使用情况",
}


def generate_word_report(
    servers: list[ServerReport],
    output_path: str | Path,
    start: str | None,
    end: str | None,
) -> Path:
    path = Path(output_path)
    document = Document()
    document.add_heading("nmon资源使用报告", level=0)
    document.add_paragraph(f"时间范围：{start or '文件全部时间'} ~ {end or '文件全部时间'}")

    _add_overview_table(document, servers)
    _add_theme_sections(document, servers)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def _add_overview_table(document: Document, servers: list[ServerReport]) -> None:
    document.add_heading("整体汇总", level=1)
    rows = build_server_overview(servers)
    if not rows:
        document.add_paragraph("没有可展示的数据。")
        return

    headers = [
        "文件名",
        "IP地址",
        "CPU使用率",
        "内存使用率",
        "SWAP使用率",
        "磁盘读速率(KB/s)",
        "磁盘写速率(KB/s)",
        "网络读速率(KB/s)",
        "网络写速率(KB/s)",
        "系统类型",
        "读取范围",
    ]
    keys = [
        "file_name",
        "ip_address",
        "cpu_usage",
        "memory_usage",
        "swap_usage",
        "disk_read",
        "disk_write",
        "network_read",
        "network_write",
        "system_type",
        "read_range",
    ]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row in rows:
        cells = table.add_row().cells
        for index, key in enumerate(keys):
            cells[index].text = row[key]


def _add_theme_sections(document: Document, servers: list[ServerReport]) -> None:
    themes = _ordered_themes(servers)
    for theme in themes:
        document.add_heading(THEME_TITLES.get(theme, f"{theme}使用情况"), level=1)
        added = False
        for server in servers:
            images = [item for item in server.images if item["theme"] == theme]
            for image in images:
                document.add_paragraph(f"{server.host} - {image.get('name', image['title'])}")
                document.add_picture(image["path"], width=Inches(6.3))
                added = True
        if not added:
            document.add_paragraph("没有可展示的图表。")


def _ordered_themes(servers: list[ServerReport]) -> list[str]:
    preferred = [
        "tps",
        "cpu",
        "memory",
        "disk",
        "network",
        "filesystem",
        "system",
        "process",
        "adapter",
        "jfs",
        "lpar",
        "unclassified",
    ]
    available = {
        image["theme"]
        for server in servers
        for image in server.images
    }
    ordered = [theme for theme in preferred if theme in available]
    ordered.extend(sorted(available - set(ordered)))
    return ordered
